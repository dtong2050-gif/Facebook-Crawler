"""Test Excel exporter — verify file structure, headers, data, edge cases."""

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from exporters.excel_exporter import MAX_ROWS_PER_SHEET, ExcelExporter
from scrapers.base import CommentData, PostData


def make_post(
    post_id: str,
    content: str = "Nội dung bài viết test",
    n_comments: int = 2,
) -> PostData:
    """Tạo PostData mẫu kèm comments cho mục đích test.

    Args:
        post_id: ID bài viết
        content: Nội dung text
        n_comments: Số comments cần tạo

    Returns:
        PostData với n_comments comments
    """
    post = PostData(
        post_id=post_id,
        url=f"https://www.facebook.com/posts/{post_id}",
        content=content,
        reactions=100,
        comments_count=n_comments,
        shares=10,
        posted_at=datetime(2024, 3, 15, 10, 30),
    )
    for i in range(n_comments):
        post.comments.append(
            CommentData(
                post_id=post_id,
                author=f"Người dùng {i}",
                content=f"Bình luận {i} cho bài {post_id}",
                reactions=5,
                commented_at=datetime(2024, 3, 15, 11, i % 60),
            )
        )
    return post


class TestExcelExporterBasic(unittest.TestCase):
    """Test các chức năng cơ bản của ExcelExporter."""

    def setUp(self) -> None:
        self.tmpdir = Path(tempfile.mkdtemp())
        self.exporter = ExcelExporter(self.tmpdir)

    def test_creates_xlsx_file(self) -> None:
        filepath = self.exporter.export([make_post("001")], "test_page")
        self.assertTrue(filepath.exists())
        self.assertEqual(filepath.suffix, ".xlsx")

    def test_filename_contains_page_name(self) -> None:
        filepath = self.exporter.export([make_post("001")], "mypage")
        self.assertIn("mypage", filepath.name)

    def test_filename_starts_with_crawl(self) -> None:
        filepath = self.exporter.export([make_post("001")], "mypage")
        self.assertTrue(filepath.name.startswith("crawl_"))

    def test_file_in_correct_directory(self) -> None:
        filepath = self.exporter.export([make_post("001")], "test")
        self.assertEqual(filepath.parent, self.tmpdir)

    def test_empty_posts_list(self) -> None:
        filepath = self.exporter.export([], "empty")
        wb = openpyxl.load_workbook(filepath)
        self.assertIn("Posts", wb.sheetnames)
        self.assertIn("Comments", wb.sheetnames)


class TestExcelExporterSheets(unittest.TestCase):
    """Test cấu trúc sheets."""

    def setUp(self) -> None:
        self.tmpdir = Path(tempfile.mkdtemp())
        self.exporter = ExcelExporter(self.tmpdir)

    def test_has_posts_sheet(self) -> None:
        filepath = self.exporter.export([make_post("001")], "test")
        wb = openpyxl.load_workbook(filepath)
        self.assertIn("Posts", wb.sheetnames)

    def test_has_comments_sheet(self) -> None:
        filepath = self.exporter.export([make_post("001")], "test")
        wb = openpyxl.load_workbook(filepath)
        self.assertIn("Comments", wb.sheetnames)

    def test_no_default_sheet(self) -> None:
        filepath = self.exporter.export([make_post("001")], "test")
        wb = openpyxl.load_workbook(filepath)
        self.assertNotIn("Sheet", wb.sheetnames)

    def test_posts_overflow_creates_second_sheet(self) -> None:
        posts = [make_post(str(i), n_comments=0) for i in range(MAX_ROWS_PER_SHEET + 3)]
        filepath = self.exporter.export(posts, "overflow")
        wb = openpyxl.load_workbook(filepath)
        self.assertIn("Posts", wb.sheetnames)
        self.assertIn("Posts_2", wb.sheetnames)

    def test_comments_overflow_creates_second_sheet(self) -> None:
        # 1 post với đủ comments để tràn sheet
        post = make_post("001", n_comments=MAX_ROWS_PER_SHEET + 5)
        filepath = self.exporter.export([post], "overflow")
        wb = openpyxl.load_workbook(filepath)
        self.assertIn("Comments", wb.sheetnames)
        self.assertIn("Comments_2", wb.sheetnames)


class TestExcelExporterHeaders(unittest.TestCase):
    """Test headers của các sheets."""

    def setUp(self) -> None:
        self.tmpdir = Path(tempfile.mkdtemp())
        self.exporter = ExcelExporter(self.tmpdir)
        self.filepath = self.exporter.export([make_post("001")], "test")
        self.wb = openpyxl.load_workbook(self.filepath)

    def test_posts_headers_correct(self) -> None:
        ws = self.wb["Posts"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 10)]
        expected = [
            "STT", "Post_ID", "Content", "Reactions",
            "Comments_Count", "Shares", "Posted_At", "Scraped_At", "URL",
        ]
        self.assertEqual(headers, expected)

    def test_comments_headers_correct(self) -> None:
        ws = self.wb["Comments"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 9)]
        expected = [
            "STT", "Post_ID", "Author", "Content",
            "Reactions", "Reply_To", "Commented_At", "Scraped_At",
        ]
        self.assertEqual(headers, expected)

    def test_header_row_is_bold(self) -> None:
        ws = self.wb["Posts"]
        self.assertTrue(ws.cell(row=1, column=1).font.bold)

    def test_freeze_panes_set(self) -> None:
        ws = self.wb["Posts"]
        self.assertEqual(ws.freeze_panes, "A2")

    def test_auto_filter_active(self) -> None:
        ws = self.wb["Posts"]
        self.assertIsNotNone(ws.auto_filter.ref)


class TestExcelExporterData(unittest.TestCase):
    """Test data được ghi đúng."""

    def setUp(self) -> None:
        self.tmpdir = Path(tempfile.mkdtemp())
        self.exporter = ExcelExporter(self.tmpdir)

    def test_posts_row_count(self) -> None:
        posts = [make_post("001"), make_post("002"), make_post("003")]
        filepath = self.exporter.export(posts, "test")
        wb = openpyxl.load_workbook(filepath)
        # Header row + 3 data rows
        self.assertEqual(wb["Posts"].max_row, 4)

    def test_comments_row_count(self) -> None:
        posts = [make_post("001", n_comments=3), make_post("002", n_comments=2)]
        filepath = self.exporter.export(posts, "test")
        wb = openpyxl.load_workbook(filepath)
        # Header + 5 comments
        self.assertEqual(wb["Comments"].max_row, 6)

    def test_stt_column_sequential(self) -> None:
        posts = [make_post("001"), make_post("002"), make_post("003")]
        filepath = self.exporter.export(posts, "test")
        wb = openpyxl.load_workbook(filepath)
        ws = wb["Posts"]
        stt_values = [ws.cell(row=r, column=1).value for r in range(2, 5)]
        self.assertEqual(stt_values, [1, 2, 3])

    def test_post_id_written_correctly(self) -> None:
        filepath = self.exporter.export([make_post("XYZ123")], "test")
        wb = openpyxl.load_workbook(filepath)
        post_id = wb["Posts"].cell(row=2, column=2).value
        self.assertEqual(post_id, "XYZ123")

    def test_unicode_content_preserved(self) -> None:
        content = "Xin chào 🇻🇳 Việt Nam ơi! Thế giới chào 你好"
        filepath = self.exporter.export([make_post("001", content=content)], "test")
        wb = openpyxl.load_workbook(filepath)
        cell_value = wb["Posts"].cell(row=2, column=3).value
        self.assertIn("Xin chào", cell_value)
        self.assertIn("🇻🇳", cell_value)

    def test_timestamp_formatted_as_string(self) -> None:
        filepath = self.exporter.export([make_post("001")], "test")
        wb = openpyxl.load_workbook(filepath)
        posted_at = wb["Posts"].cell(row=2, column=7).value
        # Phải là string, không phải datetime object
        self.assertIsInstance(posted_at, str)
        self.assertIn("2024-03-15", posted_at)

    def test_no_comments_case(self) -> None:
        post = make_post("001", n_comments=0)
        filepath = self.exporter.export([post], "test")
        wb = openpyxl.load_workbook(filepath)
        # Chỉ có header row, không có data
        self.assertEqual(wb["Comments"].max_row, 1)

    def test_comment_post_id_matches_post(self) -> None:
        filepath = self.exporter.export([make_post("POSTID42", n_comments=1)], "test")
        wb = openpyxl.load_workbook(filepath)
        comment_post_id = wb["Comments"].cell(row=2, column=2).value
        self.assertEqual(comment_post_id, "POSTID42")


if __name__ == "__main__":
    unittest.main(verbosity=2)
